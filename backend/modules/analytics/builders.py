from io import BytesIO

from docx import Document

from backend.app.logger_config import get_logger
from backend.modules.analytics.interfaces import BaseAnalysisReportBuilder

logger = get_logger(__name__)

class DocxAnalysisReportBuilder(BaseAnalysisReportBuilder):
    """Генератор DOCX-отчёта результата анализа."""

    def build(
        self,
        text: str,
    ) -> BytesIO:
        """Создаёт DOCX-файл из текста результата анализа."""

        logger.info("DOCX analysis report building started")

        try:
            buffer = BytesIO()
            document = Document()

            document.add_heading(
                "Результат анализа договора",
                level=1,
            )

            paragraphs_count = 0

            for line in text.splitlines():
                clean_line = line.strip()

                if clean_line:
                    document.add_paragraph(clean_line)
                    paragraphs_count += 1

            document.save(buffer)
            buffer.seek(0)

            logger.info(
                "DOCX analysis report built successfully: paragraphs=%s",
                paragraphs_count,
            )

            return buffer

        except Exception:
            logger.exception("Failed to build DOCX analysis report")
            raise
