from io import BytesIO

from docx import Document
from docx.text.paragraph import Paragraph
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


class AnalysisReportService:
    """Генерирует DOCX-отчёты из результата анализа."""

    def build_docx(
        self,
        text: str
    ) -> BytesIO:

        buffer = BytesIO()

        document = Document()

        document.add_heading(
            "Результат анализа договора",
            level=1
        )

        for line in text.splitlines():

            clean_line = line.strip()

            if clean_line:
                document.add_paragraph(clean_line)

        document.save(buffer)

        buffer.seek(0)

        return buffer


analysis_report_service = AnalysisReportService()