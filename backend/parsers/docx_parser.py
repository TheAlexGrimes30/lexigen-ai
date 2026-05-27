import re
from io import BytesIO

from docx import Document


class DocxParser:
    """Извлекает текст из DOCX-файла без сохранения документа на диск."""

    def extract_text_from_bytes(
        self,
        content: bytes,
    ) -> str:
        document = Document(BytesIO(content))

        parts: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = []

                for cell in row.cells:
                    text = cell.text.strip()

                    if text:
                        cells.append(text)

                if cells:
                    parts.append(" | ".join(cells))

        return self.clean_text("\n".join(parts))

    def clean_text(
        self,
        text: str,
    ) -> str:
        text = text.replace("\xa0", " ")

        text = re.sub(
            r"5\.\s*РЕКВИЗИТЫ И ПОДПИСИ СТОРОН.*",
            "",
            text,
            flags=re.DOTALL | re.IGNORECASE,
        )

        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{2,}", "\n", text)

        return text.strip()